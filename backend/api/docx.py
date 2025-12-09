"""DOCX template management endpoints."""

from __future__ import annotations

import io
import re
from typing import Any, Dict

from docx import Document
from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse

from constants import GEMINI_MODEL, TEMPLATES_DIR
from document_processing import replace_placeholders
from schemas import GenerateRequest, PolishRequest
from services import analyze_template, ensure_metadata, public_auto_sections, public_schema

router = APIRouter(prefix="", tags=["docx"])

INSTRUCTION_PATTERNS = {
	"word_count_exact": r"\b(?:write|expand|elaborate)\s+(?:to|in)\s+(\d+)\s+words?\b",
	"tone": r"\btone\s*:\s*(formal|professional|business|technical|friendly|casual)\b",
}


def extract_user_instructions(text: str):
	rules: Dict[str, Any] = {}
	cleaned = text

	match = re.search(INSTRUCTION_PATTERNS["word_count_exact"], text, flags=re.I)
	if match:
		rules["word_count_exact"] = int(match.group(1))
		cleaned = cleaned.replace(match.group(0), "")

	match = re.search(INSTRUCTION_PATTERNS["tone"], text, flags=re.I)
	if match:
		rules["tone"] = match.group(1).lower()
		cleaned = cleaned.replace(match.group(0), "")

	return cleaned.strip(), rules


POLISH_DELIMITER = r"\n?###\n?"


def clean_variation(text: str) -> str:
	return text.strip()


@router.post("/upload-template")
async def upload_template(file: UploadFile = File(...)):
	if not file.filename.lower().endswith(".docx"):
		raise HTTPException(400, "Only .docx files supported")

	template_id = file.filename.replace(" ", "_")
	path = TEMPLATES_DIR / template_id
	path.write_bytes(await file.read())

	metadata = analyze_template(template_id)
	placeholders = metadata.get("placeholders", [])
	auto_sections_internal = metadata.get("auto_sections", [])
	auto_sections_public = public_auto_sections(metadata)
	schema_public = public_schema(metadata)
	placeholder_names = (
		placeholders if placeholders else [section["id"] for section in auto_sections_internal]
	)

	return {
		"template_id": template_id,
		"placeholders": placeholder_names,
		"auto_sections": auto_sections_public,
		"schema": schema_public,
	}


@router.get("/template/{template_id}/fields")
async def template_fields(template_id: str):
	path = TEMPLATES_DIR / template_id
	if not path.exists():
		raise HTTPException(404, "Template not found")

	metadata = ensure_metadata(template_id)
	placeholders = metadata.get("placeholders", [])
	auto_sections_internal = metadata.get("auto_sections", [])
	auto_sections_public = public_auto_sections(metadata)
	schema_public = public_schema(metadata)
	placeholder_names = (
		placeholders if placeholders else [section["id"] for section in auto_sections_internal]
	)
	return {
		"placeholders": placeholder_names,
		"auto_sections": auto_sections_public,
		"schema": schema_public,
	}


@router.get("/templates")
async def list_templates():
	templates = sorted(path.name for path in TEMPLATES_DIR.glob("*.docx"))
	return {"templates": templates}


@router.post("/generate")
async def generate(req: GenerateRequest):
	path = TEMPLATES_DIR / req.template_id
	if not path.exists():
		raise HTTPException(404, "Template not found")

	metadata = ensure_metadata(req.template_id)
	payload = req.fields if req.fields is not None else req.data or {}

	doc = Document(path)
	replace_placeholders(doc, payload, metadata)

	buffer = io.BytesIO()
	doc.save(buffer)
	buffer.seek(0)

	headers = {"Content-Disposition": f'attachment; filename="{req.output_filename}"'}

	return StreamingResponse(
		buffer,
		media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
		headers=headers,
	)


@router.post("/polish")
async def polish(req: PolishRequest):
	if GEMINI_MODEL is None:
		return {"error": "Gemini API key not configured"}

	core, rules = extract_user_instructions(req.text)

	rule_lines = []
	if "word_count_exact" in rules:
		rule_lines.append(f"- EXACTLY {rules['word_count_exact']} words.")
	if "tone" in rules:
		rule_lines.append(f"- Tone: {rules['tone']}")

	rule_block = "\n".join(rule_lines)

	prompt = f"""
You are a senior technical editor. Rewrite the provided text into EXACTLY THREE alternative polished versions.
Style rules:
- Preserve the original structure: keep headings, paragraph breaks, bullet/numbered lists, and emphasis markers when they exist.
- You may improve clarity (e.g., add bold headings or consistent bullets) but keep the same sections and ordering of ideas.
- Respect ALL user instructions.
{rule_block}

Output format rules:
- Separate each alternative with a line that contains exactly three hash characters (`###`).
- Do not prepend numbers or extra labels to the alternatives.

Text to polish:
{core}
"""

	result = GEMINI_MODEL.generate_content(prompt)
	raw = result.text or ""

	lines = [clean_variation(chunk) for chunk in re.split(POLISH_DELIMITER, raw) if chunk.strip()]
	lines = lines[:3]

	return {"suggestions": lines}


__all__ = ["router"]
