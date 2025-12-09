from __future__ import annotations

from typing import Any, Dict, List

from docx import Document

from constants import SCHEMA_VERSION, TEMPLATES_DIR
from document_processing import (
    build_schema,
    extract_placeholders,
    infer_sections,
    summarize_tables,
)
from metadata_store import load_metadata, save_metadata


def analyze_template(template_id: str) -> Dict[str, Any]:
    path = TEMPLATES_DIR / template_id
    doc = Document(path)
    placeholders = extract_placeholders(doc)
    sections = infer_sections(doc)
    tables = summarize_tables(doc, sections)
    schema, field_map = build_schema(doc, sections, placeholders, tables)
    metadata = {
        "placeholders": placeholders,
        "auto_sections": sections,
        "schema": schema,
        "field_map": field_map,
        "tables": tables,
        "schema_version": SCHEMA_VERSION,
    }
    save_metadata(template_id, metadata)
    return metadata


def ensure_metadata(template_id: str) -> Dict[str, Any]:
    metadata = load_metadata(template_id)
    if (
        not metadata
        or "tables" not in metadata
        or metadata.get("schema_version") != SCHEMA_VERSION
    ):
        metadata = analyze_template(template_id)
    return metadata


def public_auto_sections(metadata: Dict[str, Any]) -> List[Dict[str, Any]]:
    results = []
    for section in metadata.get("auto_sections", []) or []:
        results.append(
            {
                "id": section.get("id"),
                "heading": section.get("heading"),
                "default_text": section.get("default_text", ""),
            }
        )
    return results


def public_schema(metadata: Dict[str, Any]) -> List[Dict[str, Any]]:
    return metadata.get("schema", []) or []
